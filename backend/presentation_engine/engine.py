# -*- coding: utf-8 -*-
"""
محرك عروض HTML — كل شريحة صفحة HTML بمقاس 1280×720 تُصدَّر كصورة HD (2x/2560×1440)
عبر Chrome headless ثم تُجمَّع في PPTX و PDF.

هويتان جاهزتان:
  "academic"  → كحلي #0F2D4A + تركوازي #20B2AA + عاجي #F8F7F2 + رمادي #5A6E7F (افتراضي)
  "dark-tech" → الهوية الداكنة سابقاً

الاستخدام:
    python engine.py deck.json [--out name.pptx] [--only-render]
"""
import sys, os, json, re, subprocess, shutil, pathlib

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
BASE = os.path.dirname(os.path.abspath(__file__))
WORKDIR = os.environ.get("PPT_STUDIO_WORKDIR") or BASE
CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"

ICONS = {
    "building": '<path d="M3 21h18"/><path d="M5 21V8l7-5 7 5v13"/><path d="M9 21v-5h6v5"/>'
                '<path d="M8 10h.01M12 10h.01M16 10h.01M8 14h.01M12 14h.01M16 14h.01"/>',
    "laptop": '<rect x="3" y="4" width="18" height="13" rx="2"/><path d="M2 21h20"/><path d="M9 9h6v5H9z"/>',
    "phone": '<rect x="7" y="2" width="10" height="20" rx="2.5"/><path d="M10.5 5h3"/>'
             '<rect x="9.5" y="15" width="5" height="3" rx="1"/>',
    "qr": '<rect x="4" y="4" width="6" height="6" rx="1"/><rect x="14" y="4" width="6" height="6" rx="1"/>'
          '<rect x="4" y="14" width="6" height="6" rx="1"/><rect x="14" y="14" width="2.5" height="2.5"/>'
          '<rect x="17.5" y="14" width="2.5" height="2.5"/><rect x="14" y="17.5" width="2.5" height="2.5"/>'
          '<rect x="17.5" y="17.5" width="2.5" height="2.5"/>',
    "bell": '<path d="M6 8a6 6 0 0 1 12 0c0 7 3 9 3 9H3s3-2 3-9"/><path d="M10.3 21a2 2 0 0 0 3.4 0"/>',
    "cloud": '<path d="M17.5 19H7a4.5 4.5 0 0 1-.42-8.98 6 6 0 0 1 11.67 1.3A3.75 3.75 0 0 1 17.5 19z"/>',
    "gear": '<circle cx="12" cy="12" r="3"/><path d="M12 2v3M12 19v3M4.9 4.9l2.1 2.1M17 17l2.1 2.1M2 12h3M19 12h3M4.9 19.1 7 17M17 7l2.1-2.1"/>',
    "headset": '<path d="M4 14v-3a8 8 0 0 1 16 0v3"/><rect x="3" y="14" width="5" height="7" rx="2"/>'
               '<rect x="16" y="14" width="5" height="7" rx="2"/>',
    "shield": '<path d="M12 22s8-3.5 8-10V5.5L12 2 4 5.5V12c0 6.5 8 10 8 10z"/><path d="M9 12l2 2 4-4"/>',
    "alert": '<path d="M12 3 2.5 19h19L12 3z"/><path d="M12 10v4"/><circle cx="12" cy="17" r="1"/>',
    "clock": '<circle cx="12" cy="12" r="9"/><path d="M12 7v5l3 2"/>',
    "target": '<circle cx="12" cy="12" r="9"/><circle cx="12" cy="12" r="5"/><circle cx="12" cy="12" r="1.2"/>',
    "check": '<path d="M20 6 9 17l-5-5"/>',
    "users": '<circle cx="9" cy="8" r="3.5"/><path d="M2.5 20c.7-3.5 3.4-5.5 6.5-5.5s5.8 2 6.5 5.5"/>'
             '<path d="M16.5 5a3.5 3.5 0 0 1 0 7M18.5 14.9c2 .7 3 2.3 3.3 4"/>',
    "trend": '<path d="M3 17l6-6 4 4 7-8"/><path d="M15 7h5v5"/>',
    "eye": '<path d="M2 12s3.5-7 10-7 10 7 10 7-3.5 7-10 7-10-7-10-7z"/><circle cx="12" cy="12" r="3"/>',
    "flag": '<path d="M4 21V4"/><path d="M4 4c5-3 10 3 16 0v9c-6 3-11-3-16 0"/>',
    "money": '<rect x="2" y="6" width="20" height="12" rx="2"/><circle cx="12" cy="12" r="3"/>'
             '<circle cx="6" cy="10" r="1" fill="currentColor"/><circle cx="18" cy="14" r="1" fill="currentColor"/>',
    "rocket": '<path d="M12 15c-2-1-3-2.5-3-5 0-4 2-7 8-8 1 6-2 8-5 8"/><path d="M9 12c-2.5 1-4 2.5-6 2 0-2 1-4 3-5"/>'
              '<path d="M11 17c.5 2.5 2.5 4 5 4 .5-2-1-3-2-4"/><circle cx="13" cy="11" r="1.4"/>',
    "doc": '<path d="M6 2h8l4 4v16H6z"/><path d="M14 2v4h4M9 12h6M9 16h6"/>',
    "mail": '<rect x="3" y="5" width="18" height="14" rx="2"/><path d="m3 7 9 6 9-6"/>',
    "star": '<path d="m12 3 2.7 5.5 6 .9-4.3 4.2 1 6-5.4-2.9-5.4 2.9 1-6L3.3 9.4l6-.9z"/>',
    "layers": '<path d="m12 3 9 5-9 5-9-5z"/><path d="m3 12 9 5 9-5"/>',
    "calendar": '<rect x="3" y="5" width="18" height="16" rx="2"/><path d="M8 3v4M16 3v4M3 10h18"/>',
    "pin": '<path d="M12 22s7-6.5 7-12a7 7 0 0 0-14 0c0 5.5 7 12 7 12z"/><circle cx="12" cy="10" r="2.5"/>',
    "lock": '<rect x="5" y="11" width="14" height="9" rx="2"/><path d="M8 11V8a4 4 0 0 1 8 0v3"/>',
    "chat": '<path d="M4 5h16v10H8l-4 4z"/>',
    "pencil": '<path d="M17 3a2.8 2.8 0 1 1 4 4L7.5 20.5 2 22l1.5-5.5z"/>',
    "handshake": '<path d="M4 13l4 3 4-5 4 5 4-3"/><path d="M2 13h2v7H2zM20 13h2v7h-2zM11 13l-1 4M13 13l1 4"/>',
}


def badge(icon, size="badge"):
    p = ICONS.get(icon, ICONS["check"])
    return (f'<span class="{size}"><svg viewBox="0 0 24 24" fill="none" stroke="currentColor" '
            f'stroke-width="1.8" stroke-linecap="round" stroke-linejoin="round">{p}</svg></span>')


# ==================================================== THEME: academic
ACAD_CSS = """
:root{color-scheme:light}
@font-face{font-family:'Cairo Fe';src:url('../../fonts/Cairo-VF.ttf') format('truetype');font-weight:200 1000;font-display:swap}
@font-face{font-family:'Changa Fe';src:url('../../fonts/Changa-VF.ttf') format('truetype');font-weight:200 800;font-display:swap}
*{margin:0;padding:0;box-sizing:border-box}
html,body{width:1280px;height:720px;overflow:hidden}
body{font-family:'Cairo Fe','Segoe UI','Tahoma',sans-serif}
.slide{--navy:#0F2D4A;--teal:#20B2AA;--bg:#F8F7F2;--bg2:#F1F4F8;--card:#FFFFFF;--gray:#5A6E7F;--line:#E3E8EE;
  --glowA:rgba(32,178,170,.13);--glowB:rgba(15,45,74,.06);--fadeA:rgba(32,178,170,0);--wm:rgba(15,45,74,.04);
  --shadow:rgba(15,45,74,.06);--badge-bg:rgba(32,178,170,.12);--navy-badge-bg:rgba(15,45,74,.08);--badge-soft-bg:rgba(32,178,170,.1);
  --line2:#EFF2F6;--text2:#3A4A5A;--text3:#E7EEF6;--note-bg:rgba(32,178,170,.07);--note-line:rgba(32,178,170,.25);
  --row:#FBFCFD;--track:#E9EDF2;--fill-gray:#AEBBC7;--teal-mid:rgba(32,178,170,.22);--dot-ring:rgba(32,178,170,.15);
  --chip-line:rgba(32,178,170,.5);--chip-bg:rgba(255,255,255,.06);--illu-shadow:rgba(15,45,74,.12);--art-tint:rgba(248,247,242,.75);
  --fh:'Changa Fe','Segoe UI',sans-serif;--fb:'Cairo Fe','Segoe UI',sans-serif;
  position:relative;width:1280px;height:720px;padding:44px 80px 30px;overflow:hidden;color:var(--navy);
  background:
   radial-gradient(circle at 92% -8%, var(--glowA), transparent 40%),
   radial-gradient(circle at -4% 108%, var(--glowB), transparent 42%),
   var(--bg)}
.slide::after{content:"";position:absolute;left:80px;right:80px;top:0;height:5px;
  background:linear-gradient(90deg,var(--teal),var(--fadeA));border-radius:0 0 6px 6px}
.head{margin-bottom:22px}
.kicker{display:inline-flex;align-items:center;gap:10px;color:var(--teal);font-size:15px;font-weight:800;letter-spacing:2px;margin-bottom:10px}
.kicker::before{content:"";width:26px;height:3px;background:var(--teal);border-radius:2px}
.title{font-size:46px;font-weight:800;line-height:1.15;letter-spacing:-.5px;color:var(--navy)}
.uline{width:56px;height:5px;border-radius:3px;background:var(--teal);margin-top:14px}
.lead{font-size:19px;line-height:1.6;color:var(--gray);max-width:940px;margin-top:14px}
.watermark{position:absolute;left:22px;bottom:64px;font-size:220px;font-weight:800;color:var(--wm);line-height:1;font-family:Georgia,serif;pointer-events:none}
/* شريط الخلاصة السفلي */
.anchor{position:absolute;left:0;right:0;bottom:0;height:58px;background:var(--navy);display:flex;align-items:center;gap:16px;padding:0 80px;z-index:2}
.anchor .lab{color:var(--teal);font-size:13px;font-weight:800;letter-spacing:1.5px}
.anchor .txt{color:#F4F7FB;font-size:16.5px;font-weight:700}
.anchor .page{margin-right:auto;color:rgba(255,255,255,.5);font-weight:700;font-size:16px}
/* بطاقات */
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(210px,1fr));gap:16px;margin-top:24px}
.card{background:var(--card);border:1px solid var(--line);border-radius:16px;padding:20px 18px;
  box-shadow:0 6px 16px var(--shadow);display:flex;flex-direction:column;gap:12px}
.card h3{font-size:18px;font-weight:800;color:var(--navy)}
.card p{font-size:14.5px;line-height:1.55;color:var(--gray)}
.badge{width:46px;height:46px;border-radius:50%;flex:none;display:grid;place-items:center;
  background:var(--badge-bg);color:var(--teal)}
.badge svg{width:23px;height:23px}
.badge.navy{background:var(--navy-badge-bg);color:var(--navy)}
.badge.soft{width:34px;height:34px;background:var(--badge-soft-bg);color:var(--teal)}
.badge.soft svg{width:17px;height:17px}
/* صفوف أيقونية */
.rows{margin-top:20px}
.row{display:flex;gap:16px;align-items:flex-start;padding:11px 0;border-bottom:1px solid var(--line2)}
.rt h4{font-size:18px;font-weight:800;color:var(--navy)}
.rt p{font-size:14.5px;color:var(--gray);line-height:1.5;margin-top:2px}
/* عمودان */
.grid2{display:grid;grid-template-columns:1fr 1fr;gap:20px;margin-top:22px}
.panel{background:var(--card);border:1px solid var(--line);border-radius:18px;padding:20px 22px;
  box-shadow:0 6px 16px var(--shadow)}
.panel .h{display:flex;gap:12px;align-items:center;margin-bottom:12px}
.panel .h h3{font-size:18.5px;font-weight:800;color:var(--navy)}
.panel ul{list-style:none}
.panel ul li{position:relative;padding:7px 22px 7px 0;font-size:15.5px;line-height:1.5;color:var(--text2)}
.panel ul li::before{content:"";position:absolute;right:0;top:14px;width:8px;height:8px;border-radius:50%;background:var(--teal)}
/* إحصائيات */
.stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(200px,1fr));gap:16px;margin-top:24px}
.stat{background:var(--card);border:1px solid var(--line);border-radius:16px;padding:20px 18px;text-align:center;
  box-shadow:0 6px 16px var(--shadow)}
.stat .v{font-family:Georgia,'IBM Plex Sans Arabic',serif;font-size:44px;font-weight:800;color:var(--navy);line-height:1.05}
.stat .v em{color:var(--teal);font-style:normal}
.stat .l{font-size:14.5px;color:var(--gray);margin-top:8px}
.note{font-size:14.5px;color:var(--gray);line-height:1.6;margin-top:18px;padding:14px 16px;border-radius:12px;background:var(--note-bg);border:1px solid var(--note-line)}
/* جدول */
table.tbl{width:100%;border-collapse:separate;border-spacing:0;margin-top:24px;font-size:15.5px;overflow:hidden;border-radius:12px}
table.tbl th{background:var(--navy);color:#fff;padding:12px 14px;text-align:right;font-weight:700}
table.tbl td{padding:11px 14px;border-bottom:1px solid var(--line2);color:var(--text2)}
table.tbl tr:nth-child(even) td{background:var(--row)}
table.tbl td.hi{color:var(--teal);font-weight:800}
/* مخطط أعمدة أفقي */
.chart{display:flex;flex-direction:column;gap:15px;margin-top:22px}
.brow{display:grid;grid-template-columns:230px 1fr auto;gap:14px;align-items:center}
.lbl{font-size:16px;font-weight:700;color:var(--navy)}
.track{height:30px;background:var(--track);border-radius:8px;overflow:hidden}
.fill{height:100%;border-radius:8px;background:var(--teal)}
.fill.navy{background:var(--navy)}
.fill.gray{background:var(--fill-gray)}
.pct{font-weight:800;color:var(--teal);font-size:15px;min-width:44px;text-align:left}
/* مسار زمني */
.timeline{display:flex;justify-content:space-between;align-items:flex-start;gap:6px;margin-top:30px;position:relative}
.timeline::before{content:"";position:absolute;top:26px;right:8%;left:8%;height:3px;border-radius:2px;
  background:linear-gradient(90deg,var(--teal),var(--teal-mid))}
.step{position:relative;flex:1;text-align:center}
.step .dot{width:52px;height:52px;border-radius:50%;background:var(--navy);color:#fff;display:grid;place-items:center;
  font-weight:800;font-size:19px;margin:0 auto;position:relative;z-index:1;box-shadow:0 0 0 5px var(--dot-ring)}
.step .t{color:var(--navy);font-weight:800;font-size:15.5px;margin-top:12px}
.step .d{color:var(--gray);font-size:12px;margin-top:2px}
/* جدول أعمال مرقّم */
.agenda{display:flex;flex-direction:column;gap:14px;margin-top:26px;max-width:1030px}
.agenda .it{display:flex;gap:18px;align-items:flex-start;padding:16px 20px;border:1px solid var(--line);border-radius:16px;background:var(--card);box-shadow:0 4px 12px var(--shadow)}
.agenda .n{width:44px;height:44px;border-radius:50%;background:var(--navy);color:#fff;font-weight:800;font-size:18px;display:grid;place-items:center;flex:none}
.agenda .t{font-size:18px;font-weight:800;color:var(--navy)}
.agenda .d{font-size:14.5px;color:var(--gray);line-height:1.5;margin-top:2px}
/* اقتباس */
.quote-slide{display:flex;flex-direction:column;justify-content:center;height:76%}
.qt{font-size:34px;font-weight:800;line-height:1.6;color:var(--navy);position:relative;padding-right:28px}
.qt::before{content:"";position:absolute;right:0;top:6px;bottom:6px;width:6px;border-radius:3px;background:var(--teal)}
.qt em{color:var(--teal);font-style:normal}
.qt-author{margin-top:20px;font-size:17px;font-weight:800;color:var(--teal)}
.qt-role{font-size:14.5px;color:var(--gray)}
/* مقارنة ثنائية */
.cmp{display:grid;grid-template-columns:1fr 1fr;gap:20px;margin-top:24px}
.cmp .side{border-radius:18px;padding:20px 22px;border:1px solid var(--line);background:var(--card);box-shadow:0 6px 16px var(--shadow)}
.cmp .side.b{background:var(--navy);border-color:var(--navy);color:#fff}
.cmp .side h3{font-size:19px;font-weight:800;margin-bottom:6px}
.cmp .side.b h3{color:#fff}
.cmp ul{list-style:none;margin-top:10px}
.cmp li{position:relative;padding:7px 24px 7px 0;font-size:15.5px;line-height:1.5;color:var(--text2)}
.cmp .side.b li{color:#E7EEF6}
.cmp li::before{content:"+";position:absolute;right:2px;top:5px;font-weight:900;color:var(--teal)}
.cmp .side.b li::before{color:var(--teal)}
/* غلاف */
.cover .title{font-size:62px}
.cover .subtitle{font-size:30px;font-weight:800;color:var(--teal);margin:6px 0 14px}
.chips{display:flex;flex-wrap:wrap;gap:10px;margin-top:20px}
.chip{padding:8px 16px;border-radius:999px;background:var(--card);border:1px solid var(--chip-line);color:var(--navy);font-size:14.5px;font-weight:600}
.cover .team{margin-top:26px;padding-top:18px;border-top:1px solid var(--line)}
.cover .team .t{color:var(--teal);font-weight:800;font-size:15px;margin-bottom:6px}
.cover .team .n{color:var(--gray);font-size:15px}
.illu{position:absolute;left:70px;top:110px;width:470px;filter:drop-shadow(0 10px 18px var(--illu-shadow))}
.slide .illu{stroke:var(--navy)}
.slide .illu .ig{stroke:var(--teal)}
.slide .illu .ig .in{stroke:var(--navy);fill:var(--navy)}
.slide .illu .ig path:not(.in){fill:var(--teal)}
/* خاتمة */
.closing{display:flex;flex-direction:column;justify-content:center;height:100%}
.closing .kicker{color:var(--teal)}
.closing .title{font-size:60px}
.bigline{font-size:27px;font-weight:800;line-height:1.55;color:var(--navy);margin-top:20px}
.bigline em{color:var(--teal);font-style:normal}
.closing2{background:var(--navy)}
.closing2 .kicker{color:var(--teal)}
.closing2 .title{color:#fff}
.closing2 .bigline{color:#F4F7FB}
.closing2 .chip{background:var(--chip-bg);border-color:var(--chip-line);color:var(--text3)}
/* خلفية مصوّرة (هجين ذكي) */
.slide.hasart{background:transparent;background-size:cover;background-position:center}
.slide.hasart::before{content:"";position:absolute;inset:0;background:var(--art-tint);z-index:0}
.slide.hasart > *:not(.anchor){position:relative;z-index:1}
/* خطوط الهوية: Changa للعناوين والأرقام، Cairo للنصوص */
.kicker,.title,.subtitle,.card h3,.rt h4,.panel .h h3,.stat .v,
table.tbl th,.lbl,.pct,.step .dot,.step .t,.anchor .lab,.anchor .txt,.watermark{font-family:var(--fh)}
.lead,li,.rt p,.stat .l,.note,.team .n,.bigline{font-family:var(--fb)}
"""


def acad_cover(d):
    illu = "" if d.get("art") else f"""
  <svg class="illu" viewBox="0 0 470 400" fill="none" stroke-width="2" stroke-linecap="round"><g class="ig" stroke-width="2">
    <rect x="18" y="250" width="150" height="130" rx="6"/><path d="M18 282l75-52 75 52"/><circle cx="93" cy="254" r="16"/>
    <path d="M52 340v-44h82v44M52 300h16M78 300h16M104 300h16M130 300h8"/>
  </g><g class="ig" stroke-width="2">
    <rect x="205" y="120" width="235" height="150" rx="10"/><path d="M215 270V140h215v130z"/><rect x="215" y="140" width="215" height="130" rx="6"/>
    <rect x="235" y="160" width="90" height="60" rx="4"/><rect x="340" y="160" width="70" height="60" rx="4"/><path d="M235 240l28-24 20 16 30-30 42 38"/>
    <rect x="215" y="270" width="215" height="18" rx="4"/>
  </g><g class="ig" stroke-width="2">
    <rect x="330" y="295" width="80" height="95" rx="12"/><rect class="in" x="342" y="310" width="56" height="56" rx="4"/>
    <path class="in" d="M342 310h14v14h-14z"/><path class="in" d="M368 310h14v14h-14z"/>
    <path class="in" d="M370 352h14v14h-14z"/><path class="in" d="M388 352h10v14h-10z"/>
    <path d="M342 352h14v14h-14z"/><path d="M368 352h14v14h-14z"/>
  </g></svg>"""
    return f"""
<section class="slide cover">
  <div class="head" style="margin-top:70px">
    <span class="kicker">{d.get('kicker','عرض تقديمي')}</span>
    <h1 class="title">{d.get('title','')}</h1>
    <div class="subtitle">{d.get('subtitle','')}</div>
    <div class="uline"></div>
  </div>
  <div class="lead" style="max-width:680px">{d.get('lead','')}</div>
  <div class="chips" style="max-width:680px">{''.join(f'<span class="chip">{c}</span>' for c in d.get('chips',[]))}</div>
  <div class="team" style="max-width:680px">
    <div class="t">فريق العمل</div>
    <div class="n">{'  •  '.join(d.get('team',[]))}</div>
  </div>
  {illu}
  <div class="anchor"><span class="lab">الخلاصة</span><span class="txt">{d.get('takeaway','')}</span><span class="page">{d.get('num','1')}</span></div>
</section>"""


def acad_content(d):
    body = ""
    if d.get("cards"):
        body += '<div class="cards">' + ''.join(
            f'<div class="card">{badge(c.get("icon", "check"))}<h3>{c["t"]}</h3><p>{c["d"]}</p></div>'
            for c in d["cards"]) + '</div>'
    if d.get("bullets"):
        body += '<div class="rows">' + ''.join(
            f'<div class="row">{badge(b.get("icon", "check"), "badge soft")}'
            f'<div class="rt"><h4>{b.get("t","")}</h4>'
            + (f'<p>{b["d"]}</p>' if "d" in b else '') + '</div></div>' for b in d["bullets"]) + '</div>'
    return f"""
<section class="slide">
  <div class="watermark">{d.get('num','')}</div>
  <div class="head"><span class="kicker">{d.get('kicker','')}</span>
    <h1 class="title">{d.get('title','')}</h1><div class="uline"></div></div>
  {f'<div class="lead">{d["lead"]}</div>' if d.get('lead') else ''}
  {body}
  <div class="anchor"><span class="lab">الخلاصة</span><span class="txt">{d.get('takeaway','')}</span><span class="page">{d.get('num','')}</span></div>
</section>"""


def acad_twocol(d):
    def panel(h, icon, items):
        lis = ''.join(f'<li>{i}</li>' for i in items)
        return f'<div class="panel"><div class="h">{badge(icon)}<h3>{h}</h3></div><ul>{lis}</ul></div>'
    return f"""
<section class="slide">
  <div class="watermark">{d.get('num','')}</div>
  <div class="head"><span class="kicker">{d.get('kicker','')}</span>
    <h1 class="title">{d.get('title','')}</h1><div class="uline"></div></div>
  <div class="grid2">{panel(d.get('col1_t',''), d.get('col1_i','layers'), d.get('col1',[]))}{panel(d.get('col2_t',''), d.get('col2_i','rocket'), d.get('col2',[]))}</div>
  {f'<div class="note">{d["note"]}</div>' if d.get('note') else ''}
  <div class="anchor"><span class="lab">الخلاصة</span><span class="txt">{d.get('takeaway','')}</span><span class="page">{d.get('num','')}</span></div>
</section>"""


def acad_stats(d):
    stats = ''.join(
        f'<div class="stat"><div class="v">{s.get("v","")}</div><div class="l">{s.get("l","")}</div></div>'
        for s in d.get("stats", []))
    chips = ''
    if d.get("chips"):
        chips = '<div class="chips">' + ''.join(f'<span class="chip">{c}</span>' for c in d["chips"]) + '</div>'
    return f"""
<section class="slide">
  <div class="watermark">{d.get('num','')}</div>
  <div class="head"><span class="kicker">{d.get('kicker','')}</span>
    <h1 class="title">{d.get('title','')}</h1><div class="uline"></div></div>
  <div class="stats">{stats}</div>
  {chips}
  {f'<div class="note">{d["note"]}</div>' if d.get('note') else ''}
  <div class="anchor"><span class="lab">الخلاصة</span><span class="txt">{d.get('takeaway','')}</span><span class="page">{d.get('num','')}</span></div>
</section>"""


def acad_table(d):
    rows = ''.join('<tr>' + ''.join(f'<td class="hi">{c}</td>' if (i == 0 and j == 1 and d.get("hl")) else f'<td>{c}</td>'
                    for j, c in enumerate(r)) + '</tr>' for i, r in enumerate(d.get("rows", [])))
    heads = ''.join(f'<th>{h}</th>' for h in d.get("headers", []))
    return f"""
<section class="slide">
  <div class="watermark">{d.get('num','')}</div>
  <div class="head"><span class="kicker">{d.get('kicker','')}</span>
    <h1 class="title">{d.get('title','')}</h1><div class="uline"></div></div>
  <table class="tbl"><thead><tr>{heads}</tr></thead><tbody>{rows}</tbody></table>
  {f'<div class="note">{d["note"]}</div>' if d.get('note') else ''}
  <div class="anchor"><span class="lab">الخلاصة</span><span class="txt">{d.get('takeaway','')}</span><span class="page">{d.get('num','')}</span></div>
</section>"""


def acad_chart(d):
    bars = ""
    for b in d.get("bars", []):
        w = max(4, min(100, b.get("v", 0)))
        cls = "fill " + b.get("cls", "")
        bars += (f'<div class="brow"><div class="lbl">{b["l"]}</div>'
                 f'<div class="track"><div class="{cls}" style="width:{w}%"></div></div>'
                 f'<div class="pct">{b.get("p", "")}</div></div>')
    return f"""
<section class="slide">
  <div class="watermark">{d.get('num','')}</div>
  <div class="head"><span class="kicker">{d.get('kicker','')}</span>
    <h1 class="title">{d.get('title','')}</h1><div class="uline"></div></div>
  {f'<div class="lead">{d["lead"]}</div>' if d.get('lead') else ''}
  <div class="chart">{bars}</div>
  {f'<div class="note">{d["note"]}</div>' if d.get('note') else ''}
  <div class="anchor"><span class="lab">الخلاصة</span><span class="txt">{d.get('takeaway','')}</span><span class="page">{d.get('num','')}</span></div>
</section>"""


def acad_timeline(d):
    steps = ''.join(
        f'<div class="step"><div class="dot">{i+1}</div><div class="t">{s.get("t","")}</div>'
        + (f'<div class="d">{s["d"]}</div>' if "d" in s else '') + '</div>'
        for i, s in enumerate(d.get("steps", [])))
    return f"""
<section class="slide">
  <div class="watermark">{d.get('num','')}</div>
  <div class="head"><span class="kicker">{d.get('kicker','')}</span>
    <h1 class="title">{d.get('title','')}</h1><div class="uline"></div></div>
  {f'<div class="lead">{d["lead"]}</div>' if d.get('lead') else ''}
  <div class="timeline">{steps}</div>
  {f'<div class="note">{d["note"]}</div>' if d.get('note') else ''}
  <div class="anchor"><span class="lab">الخلاصة</span><span class="txt">{d.get('takeaway','')}</span><span class="page">{d.get('num','')}</span></div>
</section>"""


def acad_agenda(d):
    items = ''.join(
        f'<div class="it"><div class="n">{i+1}</div><div><div class="t">{s.get("t","")}</div>'
        + (f'<div class="d">{s["d"]}</div>' if "d" in s else '') + '</div></div>'
        for i, s in enumerate(d.get("items", d.get("steps", []))))
    return f"""
<section class="slide">
  <div class="watermark">{d.get('num','')}</div>
  <div class="head"><span class="kicker">{d.get('kicker','')}</span>
    <h1 class="title">{d.get('title','')}</h1><div class="uline"></div></div>
  {f'<div class="lead">{d["lead"]}</div>' if d.get('lead') else ''}
  <div class="agenda">{items}</div>
  {f'<div class="note">{d["note"]}</div>' if d.get('note') else ''}
  <div class="anchor"><span class="lab">الخلاصة</span><span class="txt">{d.get('takeaway','')}</span><span class="page">{d.get('num','')}</span></div>
</section>"""


def acad_quote(d):
    return f"""
<section class="slide">
  <div class="watermark">{d.get('num','')}</div>
  <div class="quote-slide">
    <span class="kicker">{d.get('kicker','اقتباس')}</span>
    <div class="qt">{d.get('quote', d.get('message',''))}</div>
    <div class="qt-author">{d.get('author','')}</div>
    {f'<div class="qt-role">{d["role"]}</div>' if d.get('role') else ''}
  </div>
  <div class="anchor"><span class="lab">الخلاصة</span><span class="txt">{d.get('takeaway','')}</span><span class="page">{d.get('num','')}</span></div>
</section>"""


def acad_compare(d):
    def side(t, items, strong=False):
        lis = ''.join(f'<li>{s}</li>' for s in items)
        return f'<div class="side{" b" if strong else ""}"><h3>{t}</h3><ul>{lis}</ul></div>'
    col1 = d.get("col1", d.get("pros", []))
    col2 = d.get("col2", d.get("cons", []))
    return f"""
<section class="slide">
  <div class="watermark">{d.get('num','')}</div>
  <div class="head"><span class="kicker">{d.get('kicker','')}</span>
    <h1 class="title">{d.get('title','')}</h1><div class="uline"></div></div>
  <div class="cmp">{side(d.get('col1_t',''), col1)}{side(d.get('col2_t',''), col2, True)}</div>
  {f'<div class="note">{d["note"]}</div>' if d.get('note') else ''}
  <div class="anchor"><span class="lab">الخلاصة</span><span class="txt">{d.get('takeaway','')}</span><span class="page">{d.get('num','')}</span></div>
</section>"""


def acad_closing(d):
    return f"""
<section class="slide closing2 closing">
  <span class="kicker">{d.get('kicker','الخاتمة')}</span>
  <h1 class="title">{d.get('title','شكراً لكم')}</h1>
  <div class="bigline">{d.get('message','')}</div>
  <div class="chips">{''.join(f'<span class="chip">{c}</span>' for c in d.get('chips',[]))}</div>
  <div class="anchor" style="background:var(--navy)"><span class="lab">الخلاصة</span><span class="txt">{d.get('takeaway','')}</span><span class="page">{d.get('num','')}</span></div>
</section>"""


def acad_quote(d):
    return f"""
<section class="slide closing2">
  <span class="kicker">{d.get('kicker','اقتباس ملهم')}</span>
  <div class="bigline" style="font-size:34px;line-height:1.9">“{d.get('quote','')}”</div>
  <div class="chips"><span class="chip">— {d.get('author','')}</span></div>
  <div class="anchor" style="background:var(--navy)"><span class="lab">الخلاصة</span><span class="txt">{d.get('takeaway','')}</span><span class="page">{d.get('num','')}</span></div>
</section>"""


ACAD = {
    "css": ACAD_CSS,
    "templates": {
        "cover": acad_cover, "content": acad_content, "twocol": acad_twocol,
        "stats": acad_stats, "table": acad_table, "chart": acad_chart,
        "timeline": acad_timeline, "agenda": acad_agenda, "quote": acad_quote,
        "compare": acad_compare, "closing": acad_closing,
    },
}

# ==================================================== THEME: dark-tech (سابق)
DARK_CSS = """
:root{color-scheme:dark}
@font-face{font-family:'Cairo Fe';src:url('../../fonts/Cairo-VF.ttf') format('truetype');font-weight:200 1000;font-display:swap}
@font-face{font-family:'Changa Fe';src:url('../../fonts/Changa-VF.ttf') format('truetype');font-weight:200 800;font-display:swap}
*{margin:0;padding:0;box-sizing:border-box}
html,body{width:1280px;height:720px;overflow:hidden}
body{font-family:'Cairo Fe',"Segoe UI","Tahoma",sans-serif;background:var(--bgDark);color:var(--text)}
.slide{--main:#e3b341;--main-soft:#e3b3411a;--main-line:#e3b34155;--main-glow:rgba(227,179,65,.20);
  --surface2:#0e1628;--line-dim:#ffffff12;--line-dim2:#ffffff0d;--text-dim:#66738c;--text-mid:#9fa9be;
  --text-bright:#c6cfe0;--text-strong:#eef2fa;--text-chip:#dbe3f2;--wm2:#ffffff08;--grid:#ffffff05;
  --chip-bg2:#ffffff0a;--bars:#3d4b6b;--art-tint:rgba(11,18,32,.78);
  position:relative;width:1280px;height:720px;padding:64px 84px 58px;overflow:hidden;background:var(--bgDark)}
.a-sky{--main:#4cc2ff;--main-soft:#4cc2ff1a;--main-line:#4cc2ff55;--main-glow:rgba(76,194,255,.18)}
.a-purple{--main:#9b6bff;--main-soft:#9b6bff1a;--main-line:#9b6bff55;--main-glow:rgba(155,107,255,.18)}
.a-teal{--main:#3fd6c4;--main-soft:#3fd6c41a;--main-line:#3fd6c455;--main-glow:rgba(63,214,196,.16)}
.a-rose{--main:#ff7a90;--main-soft:#ff7a901a;--main-line:#ff7a9055;--main-glow:rgba(255,122,144,.16)}
.slide::before{content:"";position:absolute;inset:0;pointer-events:none;
  background:radial-gradient(circle at 82% -8%,var(--main-glow),transparent 46%),radial-gradient(circle at -6% 108%,var(--main-glow),transparent 42%),
  linear-gradient(var(--grid) 1px,transparent 1px),linear-gradient(90deg,var(--grid) 1px,transparent 1px);
  background-size:auto,auto,48px 48px,48px 48px;opacity:.7;mask-image:radial-gradient(circle at 50% 42%,#000 30%,transparent 88%)}
.slide::after{content:"";position:absolute;left:84px;right:84px;top:0;height:6px;background:linear-gradient(90deg,var(--main),transparent 72%);opacity:.9}
.watermark{position:absolute;right:26px;bottom:6px;font-size:230px;font-weight:800;color:var(--wm2);line-height:1;font-family:Georgia,serif}
.footer{display:flex;justify-content:space-between;align-items:center;margin-top:18px;border-top:1px solid var(--line-dim);padding-top:12px;color:var(--text-dim);font-size:14px}
.kicker{display:inline-flex;align-items:center;gap:10px;color:var(--main);font-size:20px;font-weight:700;margin-bottom:12px}
.title{font-size:48px;font-weight:800;line-height:1.12;margin-bottom:12px}
.lead{font-size:21px;line-height:1.55;color:var(--text-mid);max-width:920px}
.cards{display:flex;flex-wrap:wrap;gap:12px;margin-top:26px}
.card{flex:1 1 200px;min-width:215px;background:linear-gradient(160deg,var(--surface),var(--surface2));border:1px solid var(--line-dim);border-radius:16px;padding:18px}
.card h3{font-size:18px;font-weight:700;color:var(--main);margin-bottom:6px}
.card p{font-size:15.5px;color:var(--text-bright);line-height:1.5}
.grid2{display:grid;grid-template-columns:1fr 1fr;gap:22px;margin-top:24px}
.panel{background:linear-gradient(160deg,var(--surface),var(--surface2));border:1px solid var(--line-dim);border-radius:18px;padding:20px 22px}
.panel h3{font-size:19px;font-weight:800;margin-bottom:12px}
.panel ul{list-style:none}
.panel ul li{padding:7px 0;font-size:17px;color:var(--text-strong);border-bottom:1px solid var(--line-dim2)}
.stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(190px,1fr));gap:16px;margin-top:26px}
.stat{background:linear-gradient(160deg,var(--surface),var(--surface2));border:1px solid var(--line-dim);border-radius:18px;padding:20px;text-align:center}
.stat .v{font-size:40px;font-weight:800;color:var(--main)}
.stat .l{font-size:15px;color:var(--text-mid);margin-top:6px}
table.tbl{width:100%;border-collapse:separate;border-spacing:0;margin-top:22px;font-size:16px}
table.tbl th{background:var(--main);color:var(--bgDark);padding:12px;text-align:right;font-weight:800}
table.tbl td{padding:11px;border-bottom:1px solid var(--line-dim2);color:var(--text-bright)}
.chips{display:flex;flex-wrap:wrap;gap:10px;margin-top:22px}
.chip{padding:9px 16px;border-radius:999px;background:var(--chip-bg2);border:1px solid var(--main-line);color:var(--text-chip);font-size:15px}
.flow{display:flex;align-items:center;flex-wrap:wrap;gap:10px;margin-top:26px}
.flow .node{padding:10px 16px;border-radius:12px;background:var(--main-soft);border:1px solid var(--main-line);color:var(--main);font-weight:700;font-size:16px}
.agenda{display:flex;flex-direction:column;gap:12px;margin-top:26px;max-width:1030px}
.agenda .it{display:flex;gap:16px;align-items:flex-start;padding:14px 18px;border:1px solid var(--line-dim);border-radius:14px;background:linear-gradient(160deg,var(--surface),var(--surface2))}
.agenda .n{width:38px;height:38px;border-radius:10px;background:var(--main-soft);border:1px solid var(--main-line);color:var(--main);font-weight:800;font-size:17px;display:grid;place-items:center;flex:none}
.agenda .t{font-size:17px;font-weight:800;color:var(--text-strong)}
.agenda .d{font-size:14px;color:var(--text-mid);margin-top:2px}
.quote-slide{display:flex;flex-direction:column;justify-content:center;height:70%;margin-top:70px}
.qt{font-size:34px;font-weight:800;line-height:1.6;color:var(--text-strong);position:relative;padding-right:26px}
.qt::before{content:"";position:absolute;right:0;top:8px;bottom:8px;width:5px;border-radius:3px;background:var(--main)}
.qt em{color:var(--main);font-style:normal}
.qt-author{margin-top:18px;font-size:17px;font-weight:800;color:var(--main)}
.qt-role{font-size:14.5px;color:var(--text-mid)}
.cmp{display:grid;grid-template-columns:1fr 1fr;gap:18px;margin-top:24px}
.cmp .side{border-radius:16px;padding:18px 20px;border:1px solid var(--line-dim);background:linear-gradient(160deg,var(--surface),var(--surface2))}
.cmp .side h3{font-size:18px;font-weight:800;color:var(--main);margin-bottom:6px}
.cmp ul{list-style:none;margin-top:8px}
.cmp li{position:relative;padding:7px 22px 7px 0;font-size:15.5px;color:var(--text-strong);border-bottom:1px solid var(--line-dim2)}
.cmp li::before{content:"+";position:absolute;right:2px;top:5px;font-weight:900;color:var(--main)}
.cover .title{font-size:62px}
.cover .subtitle{font-size:29px;color:var(--main);font-weight:700;margin:4px 0 18px}
.mid{margin-top:110px;text-align:center}
.bigline{font-size:27px;font-weight:800;line-height:1.55;color:var(--text-strong)}
.bigline b{color:var(--main)}
.kicker,.title,.subtitle,.card h3,.panel h3,.stat .v,table.tbl th,.node{font-family:'Changa Fe',"Segoe UI",sans-serif}
"""


def dark_cover(d):
    return f"""
<section class="slide a-{d.get('accent','gold')} cover">
  <div class="watermark">{d.get('num','1')}</div>
  <div style="margin-top:96px"><span class="kicker">{d.get('kicker','عرض تقديمي')}</span>
    <h1 class="title">{d.get('title','')}</h1><div class="subtitle">{d.get('subtitle','')}</div>
    <div class="lead">{d.get('lead','')}</div></div>
  <div class="chips">{''.join(f'<span class="chip">{c}</span>' for c in d.get('chips',[]))}</div>
  <div class="footer"><span>{d.get('brand','')}</span><span>{'  •  '.join(d.get('team',[]) or [])}</span><span>{d.get('num','1')}</span></div>
</section>"""


def dark_content(d):
    body = ""
    if d.get("bullets"):
        body += ''.join(f'<div style="padding:8px 0;border-bottom:1px solid var(--line-dim2);font-size:18px"><b style="color:var(--main)">{b.get("t", b)}</b>'
                        + (f' <span style="color:var(--text-mid)">— {b["d"]}</span>' if isinstance(b, dict) and "d" in b else '') + '</div>'
                        for b in d["bullets"])
    if d.get("cards"):
        body += '<div class="cards">' + ''.join(f'<div class="card"><h3>{c["t"]}</h3><p>{c["d"]}</p></div>' for c in d["cards"]) + '</div>'
    return f"""
<section class="slide a-{d.get('accent','gold')}">
  <div class="watermark">{d.get('num','')}</div>
  <span class="kicker">{d.get('kicker','')}</span><h1 class="title">{d.get('title','')}</h1>
  {f'<div class="lead">{d["lead"]}</div>' if d.get('lead') else ''}
  {body}
  <div class="footer"><span>{d.get('brand','')}</span><span>{d.get('num','')}</span></div>
</section>"""


def dark_twocol(d):
    def panel(h, items):
        lis = ''.join(f'<li>{i}</li>' for i in items)
        return f'<div class="panel"><h3>{h}</h3><ul>{lis}</ul></div>'
    return f"""
<section class="slide a-{d.get('accent','purple')}">
  <div class="watermark">{d.get('num','')}</div>
  <span class="kicker">{d.get('kicker','')}</span><h1 class="title">{d.get('title','')}</h1>
  <div class="grid2">{panel(d.get('col1_t',''), d.get('col1',[]))}{panel(d.get('col2_t',''), d.get('col2',[]))}</div>
  <div class="footer"><span>{d.get('brand','')}</span><span>{d.get('num','')}</span></div>
</section>"""


def dark_stats(d):
    stats = ''.join(f'<div class="stat"><div class="v">{s["v"]}</div><div class="l">{s["l"]}</div></div>' for s in d.get("stats", []))
    return f"""
<section class="slide a-{d.get('accent','sky')}">
  <div class="watermark">{d.get('num','')}</div>
  <span class="kicker">{d.get('kicker','')}</span><h1 class="title">{d.get('title','')}</h1>
  <div class="stats">{stats}</div>
  {f'<div class="lead" style="margin-top:16px">{d["note"]}</div>' if d.get('note') else ''}
  <div class="footer"><span>{d.get('brand','')}</span><span>{d.get('num','')}</span></div>
</section>"""


def dark_table(d):
    rows = ''.join('<tr>' + ''.join(f'<td>{c}</td>' for c in r) + '</tr>' for r in d.get("rows", []))
    heads = ''.join(f'<th>{h}</th>' for h in d.get("headers", []))
    return f"""
<section class="slide a-{d.get('accent','teal')}">
  <div class="watermark">{d.get('num','')}</div>
  <span class="kicker">{d.get('kicker','')}</span><h1 class="title">{d.get('title','')}</h1>
  <table class="tbl"><thead><tr>{heads}</tr></thead><tbody>{rows}</tbody></table>
  {f'<div class="lead" style="margin-top:16px;font-size:16px">{d["note"]}</div>' if d.get('note') else ''}
  <div class="footer"><span>{d.get('brand','')}</span><span>{d.get('num','')}</span></div>
</section>"""


def dark_steps(d):
    nodes = ''.join(f'<span class="node">{s}</span><span style="color:var(--bars)">→</span>' for s in d.get("steps", []))
    nodes = nodes.rstrip('<span style="color:var(--bars)">→</span>')
    return f"""
<section class="slide a-{d.get('accent','rose')}">
  <div class="watermark">{d.get('num','')}</div>
  <span class="kicker">{d.get('kicker','')}</span><h1 class="title">{d.get('title','')}</h1>
  <div class="lead">{d.get('lead','')}</div><div class="flow">{nodes}</div>
  {f'<div class="lead" style="margin-top:16px;font-size:16px">{d["note"]}</div>' if d.get('note') else ''}
  <div class="footer"><span>{d.get('brand','')}</span><span>{d.get('num','')}</span></div>
</section>"""


def dark_agenda(d):
    items = ''.join(
        f'<div class="it"><div class="n">{i+1}</div><div><div class="t">{s.get("t","")}</div>'
        + (f'<div class="d">{s["d"]}</div>' if "d" in s else '') + '</div></div>'
        for i, s in enumerate(d.get("items", d.get("steps", []))))
    return f"""
<section class="slide a-{d.get('accent','gold')}">
  <div class="watermark">{d.get('num','')}</div>
  <span class="kicker">{d.get('kicker','')}</span><h1 class="title">{d.get('title','')}</h1>
  {f'<div class="lead">{d["lead"]}</div>' if d.get('lead') else ''}
  <div class="agenda">{items}</div>
  {f'<div class="lead" style="margin-top:16px;font-size:16px">{d["note"]}</div>' if d.get('note') else ''}
  <div class="footer"><span>{d.get('brand','')}</span><span>{d.get('num','')}</span></div>
</section>"""


def dark_quote(d):
    return f"""
<section class="slide a-{d.get('accent','gold')}">
  <div class="watermark">{d.get('num','')}</div>
  <div class="quote-slide">
    <span class="kicker">{d.get('kicker','اقتباس')}</span>
    <div class="qt">{d.get('quote', d.get('message',''))}</div>
    <div class="qt-author">{d.get('author','')}</div>
    {f'<div class="qt-role">{d["role"]}</div>' if d.get('role') else ''}
  </div>
  <div class="footer"><span>{d.get('brand','')}</span><span>{d.get('num','')}</span></div>
</section>"""


def dark_compare(d):
    def side(t, items, strong=False):
        lis = ''.join(f'<li>{s}</li>' for s in items)
        b = ' style="background:var(--main-soft);border-color:var(--main-line)"' if strong else ''
        return f'<div class="side"{b}><h3>{t}</h3><ul>{lis}</ul></div>'
    col1 = d.get("col1", d.get("pros", []))
    col2 = d.get("col2", d.get("cons", []))
    return f"""
<section class="slide a-{d.get('accent','purple')}">
  <div class="watermark">{d.get('num','')}</div>
  <span class="kicker">{d.get('kicker','')}</span><h1 class="title">{d.get('title','')}</h1>
  <div class="cmp">{side(d.get('col1_t',''), col1)}{side(d.get('col2_t',''), col2, True)}</div>
  {f'<div class="lead" style="margin-top:16px;font-size:16px">{d["note"]}</div>' if d.get('note') else ''}
  <div class="footer"><span>{d.get('brand','')}</span><span>{d.get('num','')}</span></div>
</section>"""


def dark_closing(d):
    return f"""
<section class="slide a-{d.get('accent','gold')} mid">
  <span class="kicker">{d.get('kicker','نهاية العرض')}</span>
  <h1 class="title">{d.get('title','شكراً لكم')}</h1>
  <div class="bigline">{d.get('message','')}</div>
  <div class="chips" style="justify-content:center">{''.join(f'<span class="chip">{c}</span>' for c in d.get('chips',[]))}</div>
</section>"""


def dark_quote(d):
    return f"""
<section class="slide a-{d.get('accent','gold')}">
  <span class="kicker">{d.get('kicker','اقتباس ملهم')}</span>
  <div class="bigline">“{d.get('quote','')}”</div>
  <div class="chips"><span class="chip">— {d.get('author','')}</span></div>
  <div class="footer"><span>{d.get('takeaway','')}</span><span>{d.get('num','')}</span></div>
</section>"""


DARK = {
    "css": DARK_CSS,
    "templates": {
        "cover": dark_cover, "content": dark_content, "twocol": dark_twocol,
        "stats": dark_stats, "table": dark_table, "steps": dark_steps,
        "agenda": dark_agenda, "quote": dark_quote, "compare": dark_compare,
        "closing": dark_closing,
    },
}

THEMES = {"academic": ACAD, "dark-tech": DARK}
DEFAULT_THEME = "academic"

_DARK_ACCENT_MAIN = {
    "gold": "#e3b341", "sky": "#4cc2ff", "purple": "#9b6bff",
    "teal": "#3fd6c4", "rose": "#ff7a90",
}


def resolve_theme(deck):
    """deck['theme'] يمكن أن يكون نصاً (legacy) أو كائناً {base,colors,fonts,accent}.
    يعيد (theme_name, css_override)."""
    t = deck.get("theme", DEFAULT_THEME)
    if isinstance(t, dict):
        base = t.get("base") or "academic"
        if base not in THEMES:
            base = "academic"
        return base, _identity_css_override(t, base)
    name = t if t in THEMES else DEFAULT_THEME
    return name, ""


def _identity_css_override(identity, base):
    nts, formed = {}, []
    colors = identity.get("colors") or {}
    fonts = identity.get("fonts") or {}
    if base == "dark-tech":
        main = colors.get("main") or _DARK_ACCENT_MAIN.get(identity.get("accent")) or "#4cc2ff"
        bgDark = colors.get("bgDark") or "#0b1220"
        surface = colors.get("surface") or (colors.get("bgDark") or "#0b1220")
        text = colors.get("text") or "#e8edf5"
        accent = identity.get("accent") or "sky"
        surface2 = _mix(surface, "#000", .35)
        formed.append(f".slide{{--main:{main};--main-soft:{main}1a;--main-line:{main}55;"
                      f"--main-glow:{_rgba(main,.18)};--bgDark:{bgDark};--surface:{surface};--text:{text};"
                      f"--surface2:{surface2};--line-dim:{_rgba(text,.08)};--line-dim2:{_rgba(text,.05)};"
                      f"--text-dim:{_rgba(text,.5)};--text-mid:{_rgba(text,.68)};--text-bright:{_rgba(text,.84)};"
                      f"--text-strong:{text};--text-chip:{text};--wm2:{_rgba(text,.04)};--grid:{_rgba(text,.03)};"
                      f"--chip-bg2:{_rgba(text,.05)};--bars:{_mix(text,'#000',.55)};color:{text};"
                      f"--art-tint:{_rgba(bgDark,.78)}}}")
        for cls, m in _DARK_ACCENT_MAIN.items():
            formed.append(f".slide.a-{cls}{{--main:{m};--main-soft:{m}1a;--main-line:{m}55;--main-glow:{_rgba(m,.18)}}}")
        art_alpha = max(0.0, min(1.0, float(identity.get("artTint", colors.get("artTint", 0.7)))))
        formed.append(f".slide{{--art-tint:{_rgba(bgDark, art_alpha)}}}")
        formed.append(f"body{{background:{bgDark}}}")
    else:
        navy = colors.get("navy") or "#0F2D4A"
        teal = colors.get("teal") or "#20B2AA"
        bg = colors.get("bg") or "#F8F7F2"
        bg2 = colors.get("bg2") or "#F1F4F8"
        card = colors.get("card") or "#FFFFFF"
        gray = colors.get("gray") or "#5A6E7F"
        line = colors.get("line") or "#E3E8EE"
        formed.append(f".slide{{--navy:{navy};--teal:{teal};--bg:{bg};--bg2:{bg2};--card:{card};"
                      f"--gray:{gray};--line:{line};color:{navy};"
                      f"--glowA:{_rgba(teal,.13)};--glowB:{_rgba(navy,.06)};--fadeA:{_rgba(teal,0)};--wm:{_rgba(navy,.04)};"
                      f"--shadow:{_rgba(navy,.06)};--badge-bg:{_rgba(teal,.12)};--navy-badge-bg:{_rgba(navy,.08)};"
                      f"--badge-soft-bg:{_rgba(teal,.1)};--line2:{_rgba(gray,.22)};--text2:{_mix(navy,gray,.45)};"
                      f"--text3:{_rgba(card,.9)};--note-bg:{_rgba(teal,.07)};--note-line:{_rgba(teal,.25)};"
                      f"--row:{_mix(card,'#000',.04)};--track:{_mix(card,'#000',.1)};--fill-gray:{_mix(gray,'#fff',.5)};"
                      f"--teal-mid:{_rgba(teal,.22)};--dot-ring:{_rgba(teal,.15)};--chip-line:{_rgba(teal,.5)};"
                      f"--chip-bg:{_rgba(card,.08)};--illu-shadow:{_rgba(navy,.12)};--art-tint:{_rgba(bg,.75)}}}")
        # outline strips & accents are hard-coded rgba teal; inject teal-mix globally
        formed.append(f".slide .outline{{border-color:{_rgba(teal,.55)}}} "
                      f".slide .uline{{background:{teal}}} "
                      f"table.tbl th{{background:{navy}}}")
        nts["--teal"] = teal
        # secondary soft strips (rgba(32,178,170,.12)) -> teal soft
        formed.append(f".slide .badge.soft{{background:{_rgba(teal,.12)};color:{teal}}} "
                      f".slide .note{{background:{_rgba(teal,.07)};border:1px solid {_rgba(teal,.25)}}}")
        # الشفافية المركّبة فوق الخلفيات المصوّرة (--art-tint) قابلة للضبط من الهوية
        art_alpha = max(0.0, min(1.0, float(identity.get("artTint", colors.get("artTint", 0.75)))))
        formed.append(f".slide{{--art-tint:{_rgba(bg, art_alpha)}}}")

    fh = fonts.get("fh")
    fb = fonts.get("fb")
    if fh:
        fh = _font_safe(fh)
    if fb:
        fb = _font_safe(fb)
    if fh or fb:
        frag = ""
        if fh:
            frag += f"table.tbl th,.lbl,.pct,.step .dot,.step .t,.anchor .lab,.anchor .txt,.watermark{{font-family:var(--fh)}}"
        if fb:
            frag += f".lead,li,.rt p,.stat .l,.note,.team .n,.bigline{{font-family:var(--fb)}}"
        formed.append(f".slide{{" + f"--fh:'{fh or 'Changa Fe'}','Segoe UI',sans-serif;--fb:'{fb or 'Cairo Fe'}','Segoe UI',sans-serif;" + "}")
        formed.append(frag)
    return "\n".join(formed)


FONT_WHITELIST = {
    "Changa Fe", "Cairo Fe", "Tajawal", "IBM Plex Sans Arabic", "Almarai",
    "Noto Sans Arabic", "Amiri", "Aref Ruqaa", "Markazi Text", "Mada",
    "Mirza", "Scheherazade New", "Lateef", "Reem Kufi", "Zain",
    "El Messiri", "Harmattan", "Baloo Bhaijaan 2", "Lalezar", "Jomhuria",
    "Montserrat", "Poppins", "Inter", "Roboto", "Playfair Display",
}


def _font_safe(name):
    import re as _re
    n = _re.sub(r"[^0-9A-Za-z\u0600-\u06FF \-]", "", str(name)).strip()
    n2 = _re.sub(r"\s+", " ", n)
    for f in FONT_WHITELIST:
        if n2.lower() == f.lower():
            return f
    return "Cairo Fe"


def _rgba(hex_color, alpha):
    h = hex_color.lstrip("#")
    if len(h) != 6:
        return hex_color
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return f"rgba({r},{g},{b},{alpha})"


def _mix(hex_from, hex_to, t):
    """اخلط لونين HEX بنسبة t (0..1) نحو hex_to."""
    try:
        a = [int(hex_from.lstrip('#')[i:i+2], 16) for i in (0, 2, 4)]
        b = [int(hex_to.lstrip('#')[i:i+2], 16) for i in (0, 2, 4)]
        c = [round(x + (y - x) * t) for x, y in zip(a, b)]
        return f"#{c[0]:02x}{c[1]:02x}{c[2]:02x}"
    except Exception:
        return hex_from


# ==================================================== بناء/تصدير/تجميع

def _art_path(art, theme_name, art_dir="art"):
    """مسار خلفية art: art/{theme}/ أولاً ثم مسار art_dir (هوية/افتراضي)."""
    themed = os.path.join(BASE, "art", theme_name, f"bg_{art}.png")
    if os.path.isfile(themed):
        return themed
    base = os.path.join(BASE, art_dir)
    legacy = os.path.join(base, f"bg_{art}.png")
    return legacy if os.path.isfile(legacy) else None


def available_fonts():
    """عائلات الخطوط المتاحة في fonts/ (اسم العائلة = اسم الملف)."""
    try:
        return sorted({p.stem for p in pathlib.Path(os.path.join(BASE, "fonts")).glob("*.ttf")})
    except Exception:
        return []


def extra_font_faces():
    """@font-face لأي خط إضافي في fonts/ (المضمّنة افتراضياً تُتجاهل)."""
    faces = []
    for stem in available_fonts():
        if stem in ("Cairo-VF", "Changa-VF"):
            continue
        faces.append(f"@font-face{{font-family:'{stem}';src:url('../../fonts/{stem}.ttf') format('truetype');font-display:swap}}")
    return "\n".join(faces)


def render_slide(d, css, theme_name, art_dir="art"):
    t = THEMES[theme_name]["templates"]
    fn = t.get(d.get("template"))
    if not fn:
        raise SystemExit(f"قالب غير معروف {d.get('template')} لهوية {theme_name}")
    html = fn(d)
    art = d.get("art")
    art_file = _art_path(art, theme_name, art_dir) if art else None
    if art_file:
        url = pathlib.Path(art_file).resolve().as_uri()
        html = html.replace('<section class="slide',
                            f'<section class="slide hasart" style="background-image:url(\'{url}\');background-size:cover;background-position:center"', 1)
    return f'<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8"><style>{css}</style></head><body>{html}</body></html>'


def identity_art_dir(deck):
    """إن كانت الهوية البصرية تحمل art_dir (خلفيات ملوّنة بها)، أعد مسارها داخل art/."""
    t = deck.get("theme")
    if isinstance(t, dict):
        aid = t.get("art_dir")
        return os.path.join("art", str(aid)) if aid else "art"
    return "art"


def build_html(deck, css, theme_name, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    paths = []
    art_dir = identity_art_dir(deck)
    for i, d in enumerate(deck["slides"], 1):
        d = dict(d)
        d.setdefault("num", str(i).zfill(2))
        d.setdefault("brand", deck.get("brand", "عرض تقديمي"))
        open(os.path.join(out_dir, f"slide_{i:02d}.html"), "w", encoding="utf-8").write(render_slide(d, css, theme_name, art_dir))
        paths.append(os.path.join(out_dir, f"slide_{i:02d}.html"))
    prev = ['<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8"><style>body{background:#0b1220;margin:0;padding:30px 0} .pp{width:433px;margin:0 auto 46px;box-shadow:0 10px 30px #000a;border-radius:8px;overflow:hidden}</style></head><body>']
    for p in paths:
        prev.append(f'<div class="pp">{open(p, encoding="utf-8").read()}</div>')
    prev.append('</body></html>')
    open(os.path.join(out_dir, "preview.html"), "w", encoding="utf-8").write("".join(prev))
    return paths


def chrome_shot(html_path, png_path):
    url = pathlib.Path(html_path).resolve().as_uri()
    cmd = [CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars",
           "--force-device-scale-factor=2", "--window-size=1280,720",
           "--virtual-time-budget=2500", f"--screenshot={png_path}", url]
    subprocess.run(cmd, check=True, capture_output=True)


def assemble(deck_name, pngs, out_base):
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    imgs = []
    for p in sorted(pngs):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(p, 0, 0, prs.slide_width, prs.slide_height)
        imgs.append(Image.open(p).convert("RGB"))
    pptx_p = out_base + ".pptx"
    prs.save(pptx_p)
    pdf_p = out_base + ".pdf"
    imgs[0].save(pdf_p, save_all=True, append_images=imgs[1:], resolution=150)
    return pptx_p, pdf_p


def main():
    deck = json.load(open(sys.argv[1], encoding="utf-8"))
    theme_name, css_override = resolve_theme(deck)
    theme = THEMES[theme_name]
    css = theme["css"] + ("\n" + extra_font_faces() if available_fonts() else "") + ("\n" + css_override if css_override else "")
    out_prefix = None
    if "--out" in sys.argv:
        out_prefix = sys.argv[sys.argv.index("--out") + 1]
        if out_prefix.lower().endswith((".pptx", ".pdf")):
            out_prefix = out_prefix[:-5]
    work = os.path.join(WORKDIR, "build")
    shutil.rmtree(work, ignore_errors=True)
    paths = build_html(deck, css, theme_name, os.path.join(work, "slides"))
    print(f"صفحات HTML: {len(paths)}")
    out_dir = os.path.join(work, "png")
    os.makedirs(out_dir, exist_ok=True)
    pngs = []
    for i, p in enumerate(paths, 1):
        png = os.path.join(out_dir, f"slide_{i:02d}.png")
        chrome_shot(p, png)
        pngs.append(png)
        print(" rendered", i)
    if "--only-render" in sys.argv:
        return
    name = out_prefix or re.sub(r"[\\/:*?\"<>|]+", "_", deck.get("name", "presentation"))
    pptx_p, pdf_p = assemble(deck.get("name"), pngs, os.path.join(WORKDIR, name))
    print("PPTX:", pptx_p)
    print("PDF :", pdf_p)


if __name__ == "__main__":
    main()